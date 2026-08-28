"""
adhesion/rapports.py

Génération de rapports exportables (Excel/PDF) à partir de données
tabulaires génériques (en-têtes + lignes). Deux fonctions génériques
réutilisées par les trois types de rapports (statistiques, événement,
journal d'audit) pour ne pas dupliquer la mise en forme trois fois.
"""

import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet


def excel_depuis_tableau(titre: str, entetes: list[str], lignes: list[list]) -> io.BytesIO:
    """
    Construit un classeur Excel (.xlsx) à une feuille, à partir d'une
    liste d'en-têtes et d'une liste de lignes (chaque ligne = une
    liste de valeurs, dans le même ordre que les en-têtes).
    """
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = titre[:31]  # Excel limite les noms de feuille à 31 caractères

    # Ligne de titre, fusionnée sur toute la largeur du tableau.
    feuille.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(entetes))
    cellule_titre = feuille.cell(row=1, column=1, value=titre)
    cellule_titre.font = Font(size=14, bold=True, color="14315D")
    cellule_titre.alignment = Alignment(horizontal="center")

    # En-têtes de colonnes, en ligne 3 (ligne 2 laissée vide pour l'aération).
    remplissage_entete = PatternFill(start_color="14315D", end_color="14315D", fill_type="solid")
    for colonne, entete in enumerate(entetes, start=1):
        cellule = feuille.cell(row=3, column=colonne, value=entete)
        cellule.font = Font(bold=True, color="FFFFFF")
        cellule.fill = remplissage_entete

    # Données, à partir de la ligne 4.
    for ligne_index, ligne in enumerate(lignes, start=4):
        for colonne_index, valeur in enumerate(ligne, start=1):
            feuille.cell(row=ligne_index, column=colonne_index, value=valeur)

    # Largeur de colonne approximative, basée sur le contenu le plus
    # long observé (en-tête ou donnée), pour éviter un texte tronqué
    # à l'ouverture du fichier.
    for colonne_index, entete in enumerate(entetes, start=1):
        longueur_max = len(str(entete))
        for ligne in lignes:
            if colonne_index - 1 < len(ligne):
                longueur_max = max(longueur_max, len(str(ligne[colonne_index - 1])))
        lettre = feuille.cell(row=3, column=colonne_index).column_letter
        feuille.column_dimensions[lettre].width = min(longueur_max + 4, 45)

    tampon = io.BytesIO()
    classeur.save(tampon)
    tampon.seek(0)
    return tampon


def pdf_depuis_tableau(titre: str, entetes: list[str], lignes: list[list]) -> io.BytesIO:
    """
    Construit un PDF (paysage, pour laisser de la place aux tableaux
    larges) avec un titre et un tableau de données.
    """
    tampon = io.BytesIO()
    document = SimpleDocTemplate(
        tampon, pagesize=landscape(A4),
        leftMargin=1.5 * cm, rightMargin=1.5 * cm, topMargin=1.5 * cm, bottomMargin=1.5 * cm,
    )
    styles = getSampleStyleSheet()
    elements = [
        Paragraph(titre, styles["Title"]),
        Spacer(1, 12),
    ]

    donnees_tableau = [entetes] + [[str(v) for v in ligne] for ligne in lignes]
    tableau = Table(donnees_tableau, repeatRows=1)  # repeatRows : l'en-tête se répète sur chaque page
    tableau.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#14315D")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F5F5")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    elements.append(tableau)

    document.build(elements)
    tampon.seek(0)
    return tampon