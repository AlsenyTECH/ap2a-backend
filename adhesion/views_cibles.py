"""
adhesion/views_cibles.py

API des cibles (étape 2) : personnes, groupes, ASC, établissements,
organisations, zones sinistrées.

Lecture : gestionnaires des cibles, des actions sociales, des formations
ou du suivi (ils doivent pouvoir choisir des cibles). Écriture :
GERER_CIBLES. Les cibles contiennent des données personnelles
(téléphone, pièce d'identité) : jamais d'accès sans permission.
"""

import io
import math

import openpyxl
from django.db import IntegrityError, transaction
from django.db.models import Count, Q
from django.http import HttpResponse
from rest_framework import serializers, status
from rest_framework.decorators import api_view, parser_classes, permission_classes
from rest_framework.parsers import FormParser, MultiPartParser
from rest_framework.response import Response

from .cibles import (
    PERSONNE, doublons_potentiels, fusionner, ids_zone_et_descendants, normaliser_telephone, normaliser_texte,
)
from .models import TYPES_CIBLE_CHOICES, Appartenance, Cible, JournalAudit, Membre, Zone
from .permissions import APermission, APermissionUneParmi, PermissionSelonMethode

PERMISSION_CIBLES = PermissionSelonMethode(
    lecture=APermissionUneParmi("GERER_CIBLES", "GERER_ACTIONS_SOCIALES", "GERER_FORMATIONS", "GERER_SUIVI"),
    ecriture=APermission("GERER_CIBLES"),
)

CHAMPS_PERSONNE = ("prenom", "sexe", "date_naissance", "numero_identification")
CHAMPS_COLLECTIF = ("sous_type", "responsable", "effectif")
TAILLE_PAGE = 25
MAX_LIGNES_IMPORT = 5000


def _journaliser(request, type_action, description):
    JournalAudit.objects.create(type_action=type_action, description=description, compte_auteur=request.user)


# =========================================================================
# SÉRIALISATION
# =========================================================================

class CibleSerializer(serializers.ModelSerializer):
    nom_complet = serializers.CharField(read_only=True)
    zone_chemin = serializers.SerializerMethodField()
    membre_numero_adherent = serializers.SerializerMethodField()
    nombre_appartenances = serializers.IntegerField(read_only=True, default=0)
    nombre_actions = serializers.IntegerField(read_only=True, default=0)
    membre = serializers.PrimaryKeyRelatedField(
        queryset=Membre.objects.all(), required=False, allow_null=True,
    )

    class Meta:
        model = Cible
        fields = [
            "id_cible", "type_cible", "nom", "prenom", "nom_complet", "sexe", "date_naissance",
            "numero_identification", "sous_type", "responsable", "effectif", "telephone", "email", "zone",
            "zone_chemin", "adresse", "latitude", "longitude", "notes", "membre", "membre_numero_adherent",
            "actif", "date_creation", "nombre_appartenances", "nombre_actions",
        ]
        read_only_fields = ["date_creation"]

    def get_zone_chemin(self, cible):
        return cible.zone.chemin() if cible.zone else None

    def get_membre_numero_adherent(self, cible):
        return cible.membre.numero_adherent if cible.membre else None

    def validate_nom(self, nom):
        nom = nom.strip()
        if not nom:
            raise serializers.ValidationError("Le nom est requis")
        return nom

    def validate_membre(self, membre):
        if membre is None:
            return None
        autre = Cible.objects.filter(membre=membre)
        if self.instance is not None:
            autre = autre.exclude(pk=self.instance.pk)
        if autre.exists():
            raise serializers.ValidationError("Ce membre est déjà lié à une autre cible")
        return membre

    def validate(self, donnees):
        type_cible = donnees.get("type_cible", getattr(self.instance, "type_cible", None))
        if self.instance is not None and "type_cible" in donnees and donnees["type_cible"] != self.instance.type_cible:
            raise serializers.ValidationError({"type_cible": "Le type d'une cible ne peut pas être changé"})

        if type_cible == PERSONNE:
            prenom = donnees.get("prenom", getattr(self.instance, "prenom", ""))
            if not (prenom or "").strip():
                raise serializers.ValidationError({"prenom": "Le prénom est requis pour une personne"})
            for champ in CHAMPS_COLLECTIF:
                donnees.pop(champ, None)
        else:
            for champ in CHAMPS_PERSONNE:
                donnees.pop(champ, None)
            if donnees.get("membre"):
                raise serializers.ValidationError({"membre": "Seule une personne peut être liée à un membre"})
        return donnees


def _resume(cible):
    return {
        "id_cible": cible.id_cible,
        "type_cible": cible.type_cible,
        "nom_complet": cible.nom_complet,
        "telephone": cible.telephone,
        "zone_chemin": cible.zone.chemin() if cible.zone else None,
        "date_naissance": cible.date_naissance,
        "sous_type": cible.sous_type,
    }


def _doublons_json(doublons):
    return [{"cible": _resume(cible), "raisons": raisons} for cible, raisons in doublons]


def _avec_compteurs(cibles):
    return cibles.select_related("zone__parent__parent__parent", "membre").annotate(
        nombre_appartenances=Count("appartenances", distinct=True) + Count("membres_collectif", distinct=True),
        nombre_actions=Count("beneficiaire_origine__participations", distinct=True) + Count("actions", distinct=True),
    )


# =========================================================================
# LISTE / CRÉATION
# =========================================================================

@api_view(["GET", "POST"])
@permission_classes([PERMISSION_CIBLES])
def vue_cibles(request):
    """
    GET /api/admin/cibles/?type=&q=&zone=&inclure_inactives=1&page=
        Liste paginée ; `zone` inclut toutes ses sous-zones.
    POST /api/admin/cibles/ - création. Si des cibles semblables existent,
        répond 409 avec la liste des doublons potentiels, sauf "forcer": true.
    """
    if request.method == "POST":
        serializer = CibleSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        if not request.data.get("forcer"):
            doublons = doublons_potentiels(serializer.validated_data)
            if doublons:
                return Response(
                    {
                        "erreur": "Une cible semblable existe déjà : vérifiez avant de créer un doublon.",
                        "code": "DOUBLONS_POTENTIELS",
                        "doublons": _doublons_json(doublons),
                    },
                    status=status.HTTP_409_CONFLICT,
                )
        cible = serializer.save(cree_par=request.user)
        _journaliser(request, "CREATION_CIBLE", f"Cible créée : {cible.nom_complet} ({cible.get_type_cible_display()})")
        return Response(CibleSerializer(cible).data, status=status.HTTP_201_CREATED)

    cibles = Cible.objects.all()
    if request.query_params.get("inclure_inactives") != "1":
        cibles = cibles.filter(actif=True)
    if request.query_params.get("type"):
        cibles = cibles.filter(type_cible=request.query_params["type"])
    if request.query_params.get("zone"):
        try:
            cibles = cibles.filter(zone_id__in=ids_zone_et_descendants(int(request.query_params["zone"])))
        except ValueError:
            return Response({"erreur": "zone invalide"}, status=status.HTTP_400_BAD_REQUEST)
    q = request.query_params.get("q", "").strip()
    if q:
        filtre = Q()
        for mot in q.split():
            mot_normalise = normaliser_texte(mot)
            filtre &= (
                Q(nom_normalise__contains=mot_normalise) | Q(prenom_normalise__contains=mot_normalise)
                | Q(sous_type__icontains=mot) | Q(responsable__icontains=mot)
                | Q(numero_identification__icontains=mot)
            )
        telephone = normaliser_telephone(q)
        if len(telephone) >= 4:
            filtre |= Q(telephone_normalise__contains=telephone)
        cibles = cibles.filter(filtre)

    total = cibles.count()
    try:
        page = max(1, int(request.query_params.get("page", 1)))
    except ValueError:
        page = 1
    debut = (page - 1) * TAILLE_PAGE
    resultats = _avec_compteurs(cibles.order_by("nom", "prenom", "id_cible"))[debut:debut + TAILLE_PAGE]
    return Response({
        "resultats": CibleSerializer(resultats, many=True).data,
        "total": total,
        "page": page,
        "pages": max(1, math.ceil(total / TAILLE_PAGE)),
    })


@api_view(["POST"])
@permission_classes([APermission("GERER_CIBLES")])
def vue_verifier_doublons(request):
    """POST /api/admin/cibles/verifier-doublons/ - mêmes champs qu'une création, rien n'est enregistré."""
    donnees = {
        "type_cible": request.data.get("type_cible"),
        "nom": request.data.get("nom"),
        "prenom": request.data.get("prenom"),
        "telephone": request.data.get("telephone"),
        "numero_identification": request.data.get("numero_identification"),
        "date_naissance": request.data.get("date_naissance") or None,
        "zone": request.data.get("zone"),
    }
    exclure = request.data.get("exclure")
    return Response({"doublons": _doublons_json(doublons_potentiels(donnees, exclure_id=exclure))})


# =========================================================================
# FICHE
# =========================================================================

def _historique(cible):
    """Actions dont la cible a bénéficié, de la plus récente à la plus ancienne."""
    evenements = [
        {
            "nature": "ACTION",
            "id": ac.action_id,
            "titre": ac.action.titre,
            "type": ac.action.type_action.libelle,
            "date": ac.date_intervention or ac.action.date_debut,
            "statut": f"{ac.action.get_statut_display()} · {ac.get_statut_display()}",
        }
        for ac in cible.actions.select_related("action__type_action")
    ]
    if cible.beneficiaire_origine_id:
        for p in cible.beneficiaire_origine.participations.select_related("action"):
            evenements.append({
                "nature": "ACTION_SOCIALE",
                "id": p.action.id_action,
                "titre": p.action.titre,
                "type": p.action.get_type_action_display(),
                "date": p.action.date_debut,
                "statut": p.get_statut_display(),
            })
    participant = getattr(cible.membre, "participant_externe", None) if cible.membre else None
    if participant is not None:
        for inscription in participant.inscriptions.select_related("cohorte__formation"):
            cohorte = inscription.cohorte
            evenements.append({
                "nature": "FORMATION",
                "id": cohorte.id_cohorte,
                "titre": cohorte.formation.titre,
                "type": f"Formation - cohorte {cohorte.code_cohorte}",
                "date": cohorte.date_debut,
                "statut": cohorte.get_statut_display(),
            })
    return sorted(evenements, key=lambda e: str(e["date"] or ""), reverse=True)


def _appartenances(cible):
    if cible.est_personne:
        liens = cible.appartenances.select_related("collectif__zone")
        autre = "collectif"
    else:
        liens = cible.membres_collectif.select_related("personne__zone")
        autre = "personne"
    return [
        {
            "id_appartenance": lien.id_appartenance,
            "cible": _resume(getattr(lien, autre)),
            "role": lien.role,
            "date_debut": lien.date_debut,
            "date_fin": lien.date_fin,
        }
        for lien in liens
    ]


def _fiche(cible):
    cible = _avec_compteurs(Cible.objects.filter(pk=cible.pk)).get()
    return {
        **CibleSerializer(cible).data,
        "appartenances": _appartenances(cible),
        "historique": _historique(cible),
        "besoins": [
            {
                "id_besoin": b.id_besoin,
                "description": b.description,
                "type_action": b.type_action_id,
                "type_action_libelle": b.type_action.libelle if b.type_action else None,
                "priorite": b.priorite,
                "statut": b.statut,
                "date_identification": b.date_identification,
                "notes": b.notes,
            }
            for b in cible.besoins.select_related("type_action")
        ],
    }


@api_view(["GET", "PATCH", "DELETE"])
@permission_classes([PERMISSION_CIBLES])
def vue_cible(request, id_cible):
    """GET (fiche complète) / PATCH / DELETE /api/admin/cibles/<id>/"""
    try:
        cible = Cible.objects.get(pk=id_cible)
    except Cible.DoesNotExist:
        return Response({"erreur": "Cible introuvable"}, status=status.HTTP_404_NOT_FOUND)

    if request.method == "GET":
        return Response(_fiche(cible))

    if request.method == "DELETE":
        if _historique(cible):
            return Response(
                {"erreur": "Cette cible a déjà bénéficié d'actions : désactivez-la plutôt que de la supprimer"},
                status=status.HTTP_409_CONFLICT,
            )
        nom = cible.nom_complet
        cible.delete()
        _journaliser(request, "SUPPRESSION_CIBLE", f"Cible supprimée : {nom}")
        return Response(status=status.HTTP_204_NO_CONTENT)

    serializer = CibleSerializer(cible, data=request.data, partial=True)
    serializer.is_valid(raise_exception=True)
    cible = serializer.save()
    _journaliser(request, "MODIFICATION_CIBLE", f"Cible modifiée : {cible.nom_complet}")
    return Response(_fiche(cible))


# =========================================================================
# APPARTENANCES
# =========================================================================

@api_view(["POST"])
@permission_classes([APermission("GERER_CIBLES")])
def vue_ajouter_appartenance(request, id_cible):
    """
    POST /api/admin/cibles/<id>/appartenances/ - {id_cible, role, date_debut, date_fin}
    Relie une personne et un collectif, quel que soit le côté depuis lequel on part.
    """
    try:
        cible = Cible.objects.get(pk=id_cible)
        autre = Cible.objects.get(pk=request.data.get("id_cible"))
    except (Cible.DoesNotExist, ValueError, TypeError):
        return Response({"erreur": "Cible introuvable"}, status=status.HTTP_404_NOT_FOUND)

    if cible.est_personne == autre.est_personne:
        return Response(
            {"erreur": "Une appartenance relie une personne à un collectif (groupe, ASC, établissement...)"},
            status=status.HTTP_400_BAD_REQUEST,
        )
    personne, collectif = (cible, autre) if cible.est_personne else (autre, cible)
    try:
        with transaction.atomic():
            appartenance = Appartenance.objects.create(
                personne=personne,
                collectif=collectif,
                role=str(request.data.get("role", "")).strip()[:100],
                date_debut=request.data.get("date_debut") or None,
                date_fin=request.data.get("date_fin") or None,
            )
    except IntegrityError:
        return Response({"erreur": "Cette appartenance existe déjà"}, status=status.HTTP_409_CONFLICT)
    _journaliser(
        request, "AJOUT_APPARTENANCE", f"{personne.nom_complet} rattaché(e) à {collectif.nom_complet}"
    )
    return Response(
        {"id_appartenance": appartenance.id_appartenance, "fiche": _fiche(cible)}, status=status.HTTP_201_CREATED,
    )


@api_view(["DELETE"])
@permission_classes([APermission("GERER_CIBLES")])
def vue_supprimer_appartenance(request, id_appartenance):
    """DELETE /api/admin/appartenances/<id>/"""
    try:
        appartenance = Appartenance.objects.select_related("personne", "collectif").get(pk=id_appartenance)
    except Appartenance.DoesNotExist:
        return Response({"erreur": "Appartenance introuvable"}, status=status.HTTP_404_NOT_FOUND)
    description = f"{appartenance.personne.nom_complet} retiré(e) de {appartenance.collectif.nom_complet}"
    appartenance.delete()
    _journaliser(request, "SUPPRESSION_APPARTENANCE", description)
    return Response(status=status.HTTP_204_NO_CONTENT)


# =========================================================================
# FUSION DE DOUBLONS
# =========================================================================

@api_view(["POST"])
@permission_classes([APermission("GERER_CIBLES")])
def vue_fusionner(request, id_cible):
    """
    POST /api/admin/cibles/<id>/fusionner/ - {id_doublon}
    Tout ce qui concerne le doublon est rattaché à cette cible, puis le
    doublon est supprimé. Irréversible : journalisé.
    """
    try:
        principale = Cible.objects.get(pk=id_cible)
        doublon = Cible.objects.get(pk=request.data.get("id_doublon"))
    except (Cible.DoesNotExist, ValueError, TypeError):
        return Response({"erreur": "Cible introuvable"}, status=status.HTTP_404_NOT_FOUND)

    description = f"Fusion : {doublon.nom_complet} (n°{doublon.id_cible}) dans {principale.nom_complet} (n°{principale.id_cible})"
    try:
        principale = fusionner(principale, doublon)
    except ValueError as e:
        return Response({"erreur": str(e)}, status=status.HTTP_400_BAD_REQUEST)
    _journaliser(request, "FUSION_CIBLES", description)
    return Response(_fiche(principale))


# =========================================================================
# IMPORT EXCEL
# =========================================================================

COLONNES_IMPORT = [
    ("type", "Type (Personne, Groupe, ASC, Établissement, Organisation, Zone sinistrée)"),
    ("nom", "Nom (de famille pour une personne)"),
    ("prenom", "Prénom (personne)"),
    ("sexe", "Sexe (M ou F)"),
    ("date_naissance", "Date de naissance (JJ/MM/AAAA)"),
    ("telephone", "Téléphone"),
    ("piece_identite", "Pièce d'identité"),
    ("sous_type", "Précision (école élémentaire, GIE...)"),
    ("responsable", "Responsable"),
    ("effectif", "Effectif"),
    ("zone", "Zone (commune, quartier... ou Région > Département > Commune)"),
    ("adresse", "Adresse"),
    ("collectif", "Collectif de rattachement (nom d'un groupe, d'une école...)"),
    ("role", "Rôle dans le collectif"),
]

TYPES_PAR_LIBELLE = {
    **{normaliser_texte(code): code for code, _ in TYPES_CIBLE_CHOICES},
    "personne": "PERSONNE", "groupe": "GROUPE", "asc": "ASC", "etablissement": "ETABLISSEMENT",
    "ecole": "ETABLISSEMENT", "organisation": "ORGANISATION", "zone sinistree": "ZONE_SINISTREE",
    "localite": "ZONE_SINISTREE", "zone": "ZONE_SINISTREE",
}


def _en_tete_vers_cle(en_tete):
    texte = normaliser_texte(str(en_tete or ""))
    for cle, libelle in COLONNES_IMPORT:
        if texte in (normaliser_texte(cle), normaliser_texte(libelle)) or texte.startswith(normaliser_texte(cle)):
            return cle
    return None


class _ResolveurZones:
    """Trouve une zone par nom ou par chemin "Région > Département > Commune" (avec cache)."""

    def __init__(self):
        self.cache = {}

    def __call__(self, texte):
        texte = str(texte or "").strip()
        if not texte:
            return None, None
        if texte in self.cache:
            return self.cache[texte]
        parties = [normaliser_texte(p) for p in texte.split(">") if p.strip()]
        candidates = [z for z in Zone.objects.select_related("parent__parent__parent") if normaliser_texte(z.nom) == parties[-1]]
        if len(parties) > 1:
            candidates = [z for z in candidates if [normaliser_texte(n) for n in z.chemin().split(" > ")][-len(parties):] == parties]
        if not candidates:
            resultat = (None, f"zone « {texte} » inconnue (ajoutez-la dans Référentiels > Zones)")
        elif len(candidates) > 1:
            # Préférer le niveau le plus fin (une commune plutôt que le département homonyme).
            ordre = {"QUARTIER": 0, "COMMUNE": 1, "DEPARTEMENT": 2, "REGION": 3}
            candidates.sort(key=lambda z: ordre[z.niveau])
            if ordre[candidates[0].niveau] == ordre[candidates[1].niveau]:
                resultat = (None, f"zone « {texte} » ambiguë : précisez « Région > Département > {texte} »")
            else:
                resultat = (candidates[0], None)
        else:
            resultat = (candidates[0], None)
        self.cache[texte] = resultat
        return resultat


def _date_import(valeur):
    import datetime
    if valeur in (None, ""):
        return None
    if isinstance(valeur, datetime.datetime):
        return valeur.date()
    if isinstance(valeur, datetime.date):
        return valeur
    texte = str(valeur).strip()
    for format_date in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.datetime.strptime(texte, format_date).date()
        except ValueError:
            continue
    raise ValueError(f"date « {texte} » illisible (format attendu JJ/MM/AAAA)")


@api_view(["POST"])
@permission_classes([APermission("GERER_CIBLES")])
@parser_classes([MultiPartParser, FormParser])
def vue_importer_cibles(request):
    """
    POST /api/admin/cibles/importer-excel/ (champ "fichier")
    Une ligne = une cible. Les doublons potentiels ne sont PAS créés : ils
    sont listés dans le compte rendu pour vérification. La colonne
    "collectif" rattache une personne à un groupe/une école (existant ou
    créé plus haut dans le même fichier).
    """
    fichier = request.FILES.get("fichier")
    if fichier is None:
        return Response({"erreur": "Aucun fichier fourni (champ « fichier »)"}, status=status.HTTP_400_BAD_REQUEST)
    try:
        classeur = openpyxl.load_workbook(fichier, read_only=True, data_only=True)
    except Exception:
        return Response({"erreur": "Fichier Excel illisible"}, status=status.HTTP_400_BAD_REQUEST)

    lignes = classeur.active.iter_rows(values_only=True)
    en_tetes = [_en_tete_vers_cle(c) for c in next(lignes, [])]
    if "nom" not in en_tetes:
        return Response(
            {"erreur": "Colonne « nom » introuvable : utilisez le modèle à télécharger"},
            status=status.HTTP_400_BAD_REQUEST,
        )

    zone_de = _ResolveurZones()
    crees, appartenances, doublons, erreurs = 0, 0, [], []
    collectifs_du_fichier = {}

    for numero, valeurs in enumerate(lignes, start=2):
        if numero - 1 > MAX_LIGNES_IMPORT:
            erreurs.append(f"Import limité à {MAX_LIGNES_IMPORT} lignes : la suite a été ignorée")
            break
        ligne = {cle: valeurs[i] for i, cle in enumerate(en_tetes) if cle and i < len(valeurs)}
        if not any(v not in (None, "") for v in ligne.values()):
            continue

        type_texte = normaliser_texte(str(ligne.get("type") or "personne"))
        type_cible = TYPES_PAR_LIBELLE.get(type_texte)
        if type_cible is None:
            erreurs.append(f"Ligne {numero} : type « {ligne.get('type')} » inconnu")
            continue

        zone, erreur_zone = zone_de(ligne.get("zone"))
        if erreur_zone:
            erreurs.append(f"Ligne {numero} : {erreur_zone}")
            continue
        try:
            date_naissance = _date_import(ligne.get("date_naissance"))
        except ValueError as e:
            erreurs.append(f"Ligne {numero} : {e}")
            continue

        sexe = normaliser_texte(str(ligne.get("sexe") or ""))
        effectif = ligne.get("effectif")
        donnees = {
            "type_cible": type_cible,
            "nom": str(ligne.get("nom") or "").strip(),
            "prenom": str(ligne.get("prenom") or "").strip(),
            "sexe": {"m": "M", "h": "M", "homme": "M", "masculin": "M", "f": "F", "femme": "F", "feminin": "F"}.get(sexe, ""),
            "date_naissance": date_naissance,
            "telephone": str(ligne.get("telephone") or "").strip(),
            "numero_identification": str(ligne.get("piece_identite") or "").strip(),
            "sous_type": str(ligne.get("sous_type") or "").strip(),
            "responsable": str(ligne.get("responsable") or "").strip(),
            "effectif": effectif if effectif not in ("", None) else None,
            "zone": zone.pk if zone else None,
            "adresse": str(ligne.get("adresse") or "").strip(),
        }
        serializer = CibleSerializer(data=donnees)
        if not serializer.is_valid():
            champ, messages = next(iter(serializer.errors.items()))
            erreurs.append(f"Ligne {numero} : {champ} - {messages[0]}")
            continue

        semblables = doublons_potentiels(serializer.validated_data)
        if semblables:
            cible_existante, raisons = semblables[0]
            doublons.append({
                "ligne": numero,
                "nom": f"{donnees['prenom']} {donnees['nom']}".strip(),
                "ressemble_a": _resume(cible_existante),
                "raisons": raisons,
            })
            cible = cible_existante if len(semblables) == 1 else None
        else:
            with transaction.atomic():
                cible = serializer.save(cree_par=request.user)
            crees += 1

        if cible is not None and not cible.est_personne:
            collectifs_du_fichier[normaliser_texte(cible.nom)] = cible

        nom_collectif = str(ligne.get("collectif") or "").strip()
        if cible is not None and cible.est_personne and nom_collectif:
            collectif = collectifs_du_fichier.get(normaliser_texte(nom_collectif))
            if collectif is None:
                trouves = list(
                    Cible.objects.exclude(type_cible=PERSONNE).filter(nom_normalise=normaliser_texte(nom_collectif))[:2]
                )
                collectif = trouves[0] if len(trouves) == 1 else None
            if collectif is None:
                erreurs.append(f"Ligne {numero} : collectif « {nom_collectif} » introuvable (personne créée sans rattachement)")
            else:
                _, cree = Appartenance.objects.get_or_create(
                    personne=cible, collectif=collectif,
                    defaults={"role": str(ligne.get("role") or "").strip()[:100]},
                )
                appartenances += int(cree)

    classeur.close()
    _journaliser(
        request, "IMPORT_CIBLES",
        f"Import Excel de cibles : {crees} créées, {len(doublons)} doublons écartés, {len(erreurs)} erreurs",
    )
    return Response({
        "crees": crees,
        "appartenances": appartenances,
        "doublons": doublons,
        "erreurs": erreurs,
    })


@api_view(["GET"])
@permission_classes([APermission("GERER_CIBLES")])
def vue_modele_import_cibles(request):
    """GET /api/admin/cibles/modele-excel/ - classeur prêt à remplir, avec exemples."""
    classeur = openpyxl.Workbook()
    feuille = classeur.active
    feuille.title = "Cibles"
    feuille.append([libelle for _, libelle in COLONNES_IMPORT])
    exemples = [
        ["Établissement", "École élémentaire Thiaroye 2", "", "", "", "338001122", "", "École élémentaire",
         "M. Diallo (directeur)", 640, "Dakar > Pikine", "", "", ""],
        ["Groupe", "GIE Jappo", "", "", "", "771112233", "", "Groupement de femmes", "Awa Ndiaye", 25,
         "Kaolack", "", "", ""],
        ["Personne", "Ndiaye", "Awa", "F", "12/03/1985", "771112233", "", "", "", "", "Kaolack", "",
         "GIE Jappo", "Présidente"],
        ["Personne", "Sow", "Moussa", "M", "", "", "", "", "", "", "", "", "École élémentaire Thiaroye 2", "Élève"],
        ["Zone sinistrée", "Quartier Médina Gounass", "", "", "", "", "", "Inondations 2026", "Délégué de quartier",
         1200, "Dakar > Guédiawaye", "", "", ""],
    ]
    for ligne in exemples:
        feuille.append(ligne)
    for colonne in feuille.columns:
        feuille.column_dimensions[colonne[0].column_letter].width = 24

    tampon = io.BytesIO()
    classeur.save(tampon)
    reponse = HttpResponse(
        tampon.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    reponse["Content-Disposition"] = 'attachment; filename="modele_import_cibles.xlsx"'
    return reponse
